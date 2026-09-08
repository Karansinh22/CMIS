"""
reports/generator.py — Phase 6: generate documents from the context store.

Every report is a *view* of data that already exists (transcript, summary,
topics, decisions, action items) — nothing is re-transcribed or re-analysed.

Formats
-------
    docx  Minutes of Meeting (python-docx)
    pptx  Slide deck: overview, topics, decisions, action items (python-pptx)
    md    Markdown minutes (also used as the plain-text/email version)

Project-level reports aggregate every completed meeting in a project.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

from sqlalchemy.orm import Session

from db.models import ActionItem, ContextEntry, Decision, Meeting, Project, Topic, TranscriptSegment

logger = logging.getLogger(__name__)

FORMATS = ("docx", "pptx", "md")


# ── Data gathering ────────────────────────────────────────────────────────────

@dataclass
class MeetingData:
    meeting: Meeting
    summary: str
    summary_type: str
    topics: List[Topic]
    decisions: List[Decision]
    action_items: List[ActionItem]
    speakers: List[str]
    duration_seconds: float
    transcript: List[TranscriptSegment] = field(default_factory=list)


def collect_meeting(db: Session, meeting: Meeting, include_transcript: bool = True) -> MeetingData:
    entry: Optional[ContextEntry] = meeting.context_entry
    segments = sorted(meeting.segments, key=lambda s: s.segment_index)
    duration = (max((s.end_time for s in segments), default=0.0) - min((s.start_time for s in segments), default=0.0))
    speakers = sorted({(s.speaker.name or s.speaker.label) for s in segments if s.speaker})
    return MeetingData(
        meeting=meeting,
        summary=(entry.summary if entry and entry.summary else ""),
        summary_type=(entry.summary_type if entry and entry.summary_type else meeting.summary_type or "balanced"),
        topics=list(entry.topics) if entry else [],
        decisions=list(entry.decisions) if entry else [],
        action_items=list(entry.action_items) if entry else [],
        speakers=speakers,
        duration_seconds=max(0.0, duration),
        transcript=segments if include_transcript else [],
    )


# ── Helpers ───────────────────────────────────────────────────────────────────

def _fmt_date(dt: Optional[datetime]) -> str:
    return dt.strftime("%d %b %Y, %H:%M") if dt else "—"


def _fmt_duration(seconds: float) -> str:
    seconds = int(seconds)
    if seconds < 60:
        return f"{seconds}s"
    m, s = divmod(seconds, 60)
    h, m = divmod(m, 60)
    return f"{h}h {m:02d}m" if h else f"{m}m {s:02d}s"


def _fmt_time(seconds: float) -> str:
    m, s = divmod(int(seconds), 60)
    return f"{m:02d}:{s:02d}"


def _strip_md(text: str) -> str:
    """Turn the summary's light Markdown into plain text lines for docx/pptx."""
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text


def _summary_blocks(summary: str) -> List[tuple]:
    """Yield (kind, text) tuples: heading | bullet | quote | para."""
    blocks = []
    for raw in (summary or "").splitlines():
        line = raw.rstrip()
        if not line.strip() or line.strip() == "---":
            continue
        if line.startswith("#"):
            blocks.append(("heading", _strip_md(line.lstrip("# ").strip())))
        elif re.match(r"^\s*[•\-]\s+", line):
            blocks.append(("bullet", _strip_md(re.sub(r"^\s*[•\-]\s+", "", line))))
        elif re.match(r"^\s*\d+\.\s+", line):
            blocks.append(("bullet", _strip_md(re.sub(r"^\s*\d+\.\s+", "", line))))
        elif line.startswith(">"):
            blocks.append(("quote", _strip_md(line.lstrip("> ").strip())))
        elif line.startswith("|"):
            continue  # tables are rendered from structured data instead
        else:
            blocks.append(("para", _strip_md(line.strip())))
    return blocks


def _safe_name(title: str) -> str:
    name = re.sub(r"[^A-Za-z0-9._ -]+", "", title).strip().replace(" ", "_")
    return name[:60] or "meeting"


# ── Markdown ──────────────────────────────────────────────────────────────────

def render_markdown(data: MeetingData, include_transcript: bool = False) -> str:
    m = data.meeting
    out: List[str] = []
    out.append(f"# Minutes of Meeting — {m.title}")
    out.append("")
    out.append(f"- **Date:** {_fmt_date(m.date)}")
    out.append(f"- **Duration:** {_fmt_duration(data.duration_seconds)}")
    out.append(f"- **Participants:** {', '.join(data.speakers) if data.speakers else '—'}")
    out.append(f"- **Source:** {'Live recording' if m.source == 'live' else 'Uploaded recording'}")
    out.append("")

    if data.summary:
        out.append("## Summary")
        out.append("")
        out.append(data.summary.strip())
        out.append("")

    if data.topics:
        out.append("## Topics discussed")
        out.append("")
        for i, t in enumerate(data.topics, 1):
            flag = " *(recurring)*" if t.is_recurring else ""
            out.append(f"{i}. **{t.title}**{flag}")
            if t.summary:
                out.append(f"   {t.summary}")
        out.append("")

    out.append("## Decisions")
    out.append("")
    if data.decisions:
        for i, d in enumerate(data.decisions, 1):
            line = f"{i}. {d.description}"
            if d.rationale:
                line += f" — *{d.rationale}*"
            out.append(line)
    else:
        out.append("_No decisions were recorded._")
    out.append("")

    out.append("## Action items")
    out.append("")
    if data.action_items:
        out.append("| # | Task | Owner | Due | Priority | Status |")
        out.append("|---|------|-------|-----|----------|--------|")
        for i, a in enumerate(data.action_items, 1):
            out.append(
                f"| {i} | {a.description} | {a.owner or 'Unassigned'} | {a.due or '—'} | "
                f"{(a.urgency or 'low').title()} | {'Done' if a.resolved else 'Open'} |"
            )
    else:
        out.append("_No action items were recorded._")
    out.append("")

    if include_transcript and data.transcript:
        out.append("## Transcript")
        out.append("")
        for s in data.transcript:
            who = (s.speaker.name or s.speaker.label) if s.speaker else "Speaker"
            out.append(f"- `{_fmt_time(s.start_time)}` **{who}:** {s.text}")
        out.append("")

    out.append(f"_Generated by CMIS on {_fmt_date(datetime.utcnow())} UTC._")
    return "\n".join(out)


# ── Word (docx) ───────────────────────────────────────────────────────────────

def render_docx(data: MeetingData, path: Path, include_transcript: bool = True) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    m = data.meeting
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("Minutes of Meeting", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sub = doc.add_paragraph()
    run = sub.add_run(m.title)
    run.bold = True
    run.font.size = Pt(14)

    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Light List"
    for k, v in [
        ("Date", _fmt_date(m.date)),
        ("Duration", _fmt_duration(data.duration_seconds)),
        ("Participants", ", ".join(data.speakers) if data.speakers else "—"),
        ("Source", "Live recording" if m.source == "live" else "Uploaded recording"),
        ("Summary depth", data.summary_type.title()),
    ]:
        row = meta.add_row().cells
        row[0].text = k
        row[1].text = v
        row[0].paragraphs[0].runs[0].bold = True

    doc.add_paragraph()

    # Summary
    doc.add_heading("1. Summary", level=1)
    if data.summary:
        for kind, text in _summary_blocks(data.summary):
            if kind == "heading":
                doc.add_heading(text, level=2)
            elif kind == "bullet":
                doc.add_paragraph(text, style="List Bullet")
            elif kind == "quote":
                p = doc.add_paragraph()
                r = p.add_run(text)
                r.italic = True
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            else:
                doc.add_paragraph(text)
    else:
        doc.add_paragraph("No summary available.")

    # Topics
    doc.add_heading("2. Topics discussed", level=1)
    if data.topics:
        for t in data.topics:
            p = doc.add_paragraph(style="List Number")
            r = p.add_run(t.title)
            r.bold = True
            if t.is_recurring:
                p.add_run("  (recurring topic)").italic = True
            if t.summary:
                doc.add_paragraph(t.summary)
    else:
        doc.add_paragraph("No topics extracted.")

    # Decisions
    doc.add_heading("3. Decisions", level=1)
    if data.decisions:
        for d in data.decisions:
            p = doc.add_paragraph(style="List Number")
            p.add_run(d.description)
            if d.rationale:
                p.add_run(f" — {d.rationale}").italic = True
    else:
        doc.add_paragraph("No decisions were recorded.")

    # Action items
    doc.add_heading("4. Action items", level=1)
    if data.action_items:
        table = doc.add_table(rows=1, cols=6)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(["#", "Task", "Owner", "Due", "Priority", "Status"]):
            hdr[i].text = h
            hdr[i].paragraphs[0].runs[0].bold = True
        for i, a in enumerate(data.action_items, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = a.description
            row[2].text = a.owner or "Unassigned"
            row[3].text = a.due or "—"
            row[4].text = (a.urgency or "low").title()
            row[5].text = "Done" if a.resolved else "Open"
    else:
        doc.add_paragraph("No action items were recorded.")

    # Transcript appendix
    if include_transcript and data.transcript:
        doc.add_page_break()
        doc.add_heading("Appendix — Transcript", level=1)
        for s in data.transcript:
            who = (s.speaker.name or s.speaker.label) if s.speaker else "Speaker"
            p = doc.add_paragraph()
            r = p.add_run(f"[{_fmt_time(s.start_time)}] {who}: ")
            r.bold = True
            r.font.size = Pt(9)
            r2 = p.add_run(s.text)
            r2.font.size = Pt(9)

    footer = doc.add_paragraph()
    fr = footer.add_run(f"Generated by CMIS on {_fmt_date(datetime.utcnow())} UTC")
    fr.italic = True
    fr.font.size = Pt(8)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


# ── PowerPoint (pptx) ─────────────────────────────────────────────────────────

def render_pptx(data: MeetingData, path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    m = data.meeting
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def title_slide(title: str, subtitle: str = "") -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title
        if subtitle and len(slide.placeholders) > 1:
            slide.placeholders[1].text = subtitle

    def bullet_slide(title: str, bullets: List[str], max_per_slide: int = 7) -> None:
        if not bullets:
            bullets = ["Nothing recorded."]
        for start in range(0, len(bullets), max_per_slide):
            chunk = bullets[start:start + max_per_slide]
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title if start == 0 else f"{title} (cont.)"
            body = slide.placeholders[1].text_frame
            body.clear()
            for i, b in enumerate(chunk):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text = b
                p.font.size = Pt(18 if len(b) < 90 else 15)

    def table_slide(title: str, header: List[str], rows: List[List[str]], per_slide: int = 8) -> None:
        if not rows:
            bullet_slide(title, [])
            return
        for start in range(0, len(rows), per_slide):
            chunk = rows[start:start + per_slide]
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = title if start == 0 else f"{title} (cont.)"
            shape = slide.shapes.add_table(
                len(chunk) + 1, len(header), Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.4) * (len(chunk) + 1)
            )
            table = shape.table
            for c, h in enumerate(header):
                table.cell(0, c).text = h
                table.cell(0, c).text_frame.paragraphs[0].font.bold = True
                table.cell(0, c).text_frame.paragraphs[0].font.size = Pt(14)
            for r, row in enumerate(chunk, 1):
                for c, val in enumerate(row):
                    table.cell(r, c).text = val
                    table.cell(r, c).text_frame.paragraphs[0].font.size = Pt(12)

    title_slide(m.title, f"{_fmt_date(m.date)}  ·  {_fmt_duration(data.duration_seconds)}  ·  "
                         f"{len(data.speakers) or 1} participant{'s' if len(data.speakers) != 1 else ''}")

    # Overview: first paragraph-ish lines of the summary
    overview = [t for k, t in _summary_blocks(data.summary) if k in ("para", "bullet")][:6]
    bullet_slide("Overview", overview or ["No summary available."])

    topics = [f"{t.title}{' (recurring)' if t.is_recurring else ''}: {t.summary}" if t.summary else t.title for t in data.topics]
    bullet_slide("Topics discussed", topics, max_per_slide=5)

    decisions = [f"{d.description}" + (f" — {d.rationale}" if d.rationale else "") for d in data.decisions]
    bullet_slide("Decisions", decisions)

    rows = [
        [a.description, a.owner or "Unassigned", a.due or "—", (a.urgency or "low").title(), "Done" if a.resolved else "Open"]
        for a in data.action_items
    ]
    table_slide("Action items", ["Task", "Owner", "Due", "Priority", "Status"], rows)

    closing = prs.slides.add_slide(prs.slide_layouts[1])
    closing.shapes.title.text = "Next steps"
    tf = closing.placeholders[1].text_frame
    tf.clear()
    open_items = [a for a in data.action_items if not a.resolved]
    tf.paragraphs[0].text = f"{len(open_items)} open action item{'s' if len(open_items) != 1 else ''}"
    p = tf.add_paragraph()
    p.text = f"{len(data.decisions)} decision{'s' if len(data.decisions) != 1 else ''} recorded"
    p = tf.add_paragraph()
    p.text = f"Generated by CMIS on {_fmt_date(datetime.utcnow())} UTC"
    p.font.size = Pt(12)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


# ── Project-level report ──────────────────────────────────────────────────────

def render_project_docx(db: Session, project: Project, path: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    meetings = sorted(
        [mt for mt in project.meetings if mt.status == "done"],
        key=lambda x: x.date or datetime.min,
    )
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(f"Project Report — {project.name}", level=0)
    meta = [f"Company: {project.company}" if project.company else None,
            f"Category: {project.category}" if project.category else None,
            f"Meetings: {len(meetings)}"]
    doc.add_paragraph("  ·  ".join(x for x in meta if x))
    if project.description:
        doc.add_paragraph(project.description)

    if project.summary and project.summary.overall_summary:
        doc.add_heading("Overall summary", level=1)
        for kind, text in _summary_blocks(project.summary.overall_summary):
            if kind == "heading":
                doc.add_heading(text, level=2)
            elif kind == "bullet":
                doc.add_paragraph(text, style="List Bullet")
            else:
                doc.add_paragraph(text)

    # Consolidated open action items
    doc.add_heading("Open action items across meetings", level=1)
    rows = []
    for mt in meetings:
        if not mt.context_entry:
            continue
        for a in mt.context_entry.action_items:
            if not a.resolved:
                rows.append((mt.title, a.description, a.owner or "Unassigned", a.due or "—", (a.urgency or "low").title()))
    if rows:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        for i, h in enumerate(["Meeting", "Task", "Owner", "Due", "Priority"]):
            table.rows[0].cells[i].text = h
            table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        for r in rows:
            cells = table.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = v
    else:
        doc.add_paragraph("No open action items.")

    # Per-meeting decisions
    doc.add_heading("Decisions by meeting", level=1)
    any_decision = False
    for mt in meetings:
        decisions = mt.context_entry.decisions if mt.context_entry else []
        if not decisions:
            continue
        any_decision = True
        doc.add_heading(f"{mt.title} — {_fmt_date(mt.date)}", level=2)
        for d in decisions:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(d.description)
            if d.rationale:
                p.add_run(f" — {d.rationale}").italic = True
    if not any_decision:
        doc.add_paragraph("No decisions recorded yet.")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


# ── Entry point ───────────────────────────────────────────────────────────────

def generate_meeting_report(db: Session, meeting: Meeting, fmt: str, out_dir: Path,
                            include_transcript: bool = True) -> Path:
    if fmt not in FORMATS:
        raise ValueError(f"Unsupported format '{fmt}'. Choose one of {', '.join(FORMATS)}.")
    data = collect_meeting(db, meeting, include_transcript=include_transcript)
    stamp = datetime.utcnow().strftime("%Y%m%d-%H%M%S")
    base = f"{_safe_name(meeting.title)}_{stamp}"
    if fmt == "docx":
        return render_docx(data, out_dir / f"{base}_MoM.docx", include_transcript=include_transcript)
    if fmt == "pptx":
        return render_pptx(data, out_dir / f"{base}_slides.pptx")
    path = out_dir / f"{base}_MoM.md"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(render_markdown(data, include_transcript=include_transcript), encoding="utf-8")
    return path
